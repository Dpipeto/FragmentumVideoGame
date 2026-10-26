"""
Utilidades para editar TesisIA.docx (y para escribir el GDD con sus mismos estilos) sin perder lo
que ya tiene: numeración automática de tablas y figuras (campos SEQ), referencias cruzadas (campos
REF a marcadores _Ref_*), comentarios y resaltados.

Las piezas de texto son listas de:
    'texto'                     texto normal
    ('i', 'texto')              cursiva
    ('b', 'texto')              negrita
    ('hl', 'texto')             resaltado en amarillo (dato pendiente)
    ('ref', 'clave')            referencia cruzada a la tabla o figura con marcador _Ref_<clave>
"""
import copy
import re

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from lxml import etree

AUTOR = 'Asistente de IA (Claude)'
INICIALES = 'IA'

ESTILOS = {
    'texto': 'Textotesis',
    'sinsangria': 'Textotesissinsangra',
    'vineta': 'Vietatesis',
    'nota': 'Notatesis',
    'leyenda_tabla': 'Leyendatabla',
    'leyenda_figura': 'Leyendafigura',
    't1': 'Ttulo1',
    't2': 'Ttulo2',
    't3': 'Ttulo3tesis',
    't4': 'Ttulo4tesis',
    'tabla': 'Textotabla',
    'referencia': 'Referenciatesis',
    'portada': 'Portadatesis',
    'preliminar': 'Ttulopreliminar',
    'glosario': 'Glosariotesis',
    'codigo': 'Cdigofuente',
}


def texto_de(e):
    """El texto visible de un elemento (sin las instrucciones de los campos)."""
    return ''.join(t.text or '' for t in e.iter(qn('w:t')))


def estilo_de(e):
    ppr = e.find(qn('w:pPr'))
    st = ppr.find(qn('w:pStyle')) if ppr is not None else None
    return st.get(qn('w:val')) if st is not None else ''


class Documento:
    def __init__(self, ruta):
        self.doc = docx.Document(ruta)
        self.body = self.doc.element.body
        ids = [int(b.get(qn('w:id'))) for b in self.doc.element.iter(qn('w:bookmarkStart'))]
        self._marcador = (max(ids) if ids else 0) + 1
        # Los marcadores _Toc de las leyendas nuevas, después de los de una actualización anterior.
        tocs = [int(m.group(1)) for b in self.doc.element.iter(qn('w:bookmarkStart'))
                for m in [re.fullmatch(r'_Toc(811\d{6})', b.get(qn('w:name')) or '')] if m]
        self._toc = max(tocs + [811000000])
        self.plantilla_tabla = None
        for t in self.body.iter(qn('w:tbl')):
            self.plantilla_tabla = t
            break

    # --- buscar ---

    def bloques(self):
        """Los hijos del cuerpo, sin las entradas de los índices (estilos TDC y tabla de ilustraciones)."""
        return [e for e in self.body.iterchildren()
                if not (estilo_de(e).startswith('TDC') or estilo_de(e) == 'Tabladeilustraciones')]

    def buscar(self, inicio, estilo=None, desde=None, contiene=False):
        """El primer párrafo del cuerpo cuyo texto empieza por <inicio> (o lo contiene)."""
        encontrados = []
        activo = desde is None
        for e in self.bloques():
            if not activo:
                activo = e is desde
                continue
            if e.tag != qn('w:p'):
                continue
            t = texto_de(e)
            if (inicio in t) if contiene else t.startswith(inicio):
                if estilo is None or estilo_de(e) == estilo:
                    encontrados.append(e)
        if not encontrados:
            raise KeyError(f'No encuentro el párrafo «{inicio}»')
        return encontrados[0]

    def tabla_tras(self, leyenda):
        """La tabla que sigue a la leyenda que empieza por <leyenda> («Tabla 9.»)."""
        p = self.buscar(leyenda)
        e = p.getnext()
        while e is not None and e.tag != qn('w:tbl'):
            e = e.getnext()
        return e

    # --- crear ---

    def _run(self, texto, cursiva=False, negrita=False, resaltado=False, tam=None):
        r = OxmlElement('w:r')
        if cursiva or negrita or resaltado or tam:
            rpr = OxmlElement('w:rPr')
            if negrita:
                rpr.append(OxmlElement('w:b'))
            if cursiva:
                rpr.append(OxmlElement('w:i'))
            if tam:
                sz = OxmlElement('w:sz')
                sz.set(qn('w:val'), str(tam))
                rpr.append(sz)
            if resaltado:
                h = OxmlElement('w:highlight')
                h.set(qn('w:val'), 'yellow')
                rpr.append(h)
            r.append(rpr)
        t = OxmlElement('w:t')
        t.text = texto
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        r.append(t)
        return r

    def _campo(self, instruccion, resultado, negrita=False):
        """Las carreras de un campo: inicio, instrucción, separador, resultado y fin."""
        def fld(tipo):
            r = OxmlElement('w:r')
            if negrita:
                rpr = OxmlElement('w:rPr')
                rpr.append(OxmlElement('w:b'))
                r.append(rpr)
            c = OxmlElement('w:fldChar')
            c.set(qn('w:fldCharType'), tipo)
            r.append(c)
            return r
        r_ins = OxmlElement('w:r')
        if negrita:
            rpr = OxmlElement('w:rPr')
            rpr.append(OxmlElement('w:b'))
            r_ins.append(rpr)
        ins = OxmlElement('w:instrText')
        ins.text = instruccion
        ins.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        r_ins.append(ins)
        return [fld('begin'), r_ins, fld('separate'), self._run(resultado, negrita=negrita), fld('end')]

    def _carreras(self, piezas, tam=None):
        if isinstance(piezas, str):
            piezas = [piezas]
        out = []
        for pz in piezas:
            if isinstance(pz, str):
                out.append(self._run(pz, tam=tam))
            elif pz[0] == 'i':
                out.append(self._run(pz[1], cursiva=True, tam=tam))
            elif pz[0] == 'b':
                out.append(self._run(pz[1], negrita=True, tam=tam))
            elif pz[0] == 'bi':
                out.append(self._run(pz[1], negrita=True, cursiva=True, tam=tam))
            elif pz[0] == 'hl':
                out.append(self._run(pz[1], resaltado=True, tam=tam))
            elif pz[0] == 'ref':
                out.extend(self._campo(f' REF _Ref_{pz[1]} \\h ', '¿?', negrita=True))
            else:
                raise ValueError(pz)
        return out

    def parrafo(self, estilo, piezas=(), centrado=False, mantener=False):
        p = OxmlElement('w:p')
        ppr = OxmlElement('w:pPr')
        st = OxmlElement('w:pStyle')
        st.set(qn('w:val'), ESTILOS.get(estilo, estilo))
        ppr.append(st)
        if mantener:
            ppr.append(OxmlElement('w:keepNext'))
        if centrado:
            jc = OxmlElement('w:jc')
            jc.set(qn('w:val'), 'center')
            ppr.append(jc)
        p.append(ppr)
        for r in self._carreras(piezas):
            p.append(r)
        return p

    def _marcadores(self, p, nombre, primero, ultimo):
        """Rodea con un marcador las carreras de <primero> a <ultimo> (incluidas)."""
        i = self._marcador
        self._marcador += 1
        bs = OxmlElement('w:bookmarkStart')
        bs.set(qn('w:id'), str(i))
        bs.set(qn('w:name'), nombre)
        be = OxmlElement('w:bookmarkEnd')
        be.set(qn('w:id'), str(i))
        primero.addprevious(bs)
        ultimo.addnext(be)

    def leyenda(self, tipo, clave, titulo):
        """«Tabla N. Título» o «Figura N. Título», con su campo SEQ y los marcadores _Ref_ y _Toc."""
        estilo = 'leyenda_tabla' if tipo == 'Tabla' else 'leyenda_figura'
        p = self.parrafo(estilo)
        r_label = self._run(f'{tipo} ', negrita=True)
        p.append(r_label)
        campo = self._campo(f' SEQ {tipo} \\* ARABIC ', '¿?', negrita=True)
        for r in campo:
            p.append(r)
        r_punto = self._run('. ', negrita=True)
        p.append(r_punto)
        r_tit = self._run(titulo, cursiva=True)
        p.append(r_tit)
        self._marcadores(p, f'_Ref_{clave}', r_label, campo[-1])
        self._toc += 1
        self._marcadores(p, f'_Toc{self._toc}', p.find(qn('w:bookmarkStart')), r_tit)
        return p

    def nota(self, piezas):
        p = self.parrafo('nota')
        p.append(self._run('Nota. ', cursiva=True))
        for r in self._carreras(piezas):
            p.append(r)
        return p

    def figura(self, ruta, ancho_cm):
        aux = self.doc.add_paragraph()
        aux.add_run().add_picture(ruta, width=Cm(ancho_cm))
        p = aux._p
        p.getparent().remove(p)
        nuevo = self.parrafo('sinsangria', centrado=True, mantener=True)
        for r in list(p.iter(qn('w:r'))):
            nuevo.append(r)
        return nuevo

    def tabla(self, encabezados, filas, anchos):
        """Una tabla con el estilo de las de la tesis. <anchos> en twips (suman ~9350)."""
        t = OxmlElement('w:tbl')
        tblpr = copy.deepcopy(self.plantilla_tabla.find(qn('w:tblPr')))
        t.append(tblpr)
        grid = OxmlElement('w:tblGrid')
        for a in anchos:
            g = OxmlElement('w:gridCol')
            g.set(qn('w:w'), str(a))
            grid.append(g)
        t.append(grid)

        def fila(celdas, cabecera=False):
            tr = OxmlElement('w:tr')
            trpr = OxmlElement('w:trPr')
            trpr.append(OxmlElement('w:cantSplit'))
            if cabecera:
                trpr.append(OxmlElement('w:tblHeader'))
            jc = OxmlElement('w:jc')
            jc.set(qn('w:val'), 'center')
            trpr.append(jc)
            tr.append(trpr)
            for a, contenido in zip(anchos, celdas):
                tc = OxmlElement('w:tc')
                tcpr = OxmlElement('w:tcPr')
                tcw = OxmlElement('w:tcW')
                tcw.set(qn('w:w'), str(a))
                tcw.set(qn('w:type'), 'dxa')
                tcpr.append(tcw)
                tc.append(tcpr)
                partes = contenido if isinstance(contenido, list) and contenido and isinstance(contenido[0], list) else [contenido]
                for parte in partes:
                    p = OxmlElement('w:p')
                    ppr = OxmlElement('w:pPr')
                    st = OxmlElement('w:pStyle')
                    st.set(qn('w:val'), ESTILOS['tabla'])
                    ppr.append(st)
                    p.append(ppr)
                    for r in self._carreras(parte, tam=18):
                        p.append(r)
                    tc.append(p)
                tr.append(tc)
            return tr

        t.append(fila(encabezados, cabecera=True))
        for f in filas:
            t.append(fila(f))
        return t

    # --- cambiar ---

    @staticmethod
    def despues(ancla, *elementos):
        """Inserta los elementos, en orden, justo después de <ancla>. Devuelve el último."""
        for e in elementos:
            ancla.addnext(e)
            ancla = e
        return ancla

    @staticmethod
    def antes(ancla, *elementos):
        for e in elementos:
            ancla.addprevious(e)
        return elementos[-1] if elementos else ancla

    def poner(self, p, piezas):
        """Cambia el texto de un párrafo, conservando su estilo, sus marcadores y sus comentarios."""
        conservar_ini, conservar_fin = [], []
        for hijo in list(p):
            if hijo.tag == qn('w:pPr'):
                continue
            if hijo.tag in (qn('w:bookmarkStart'), qn('w:commentRangeStart')):
                conservar_ini.append(hijo)
            elif hijo.tag in (qn('w:bookmarkEnd'), qn('w:commentRangeEnd')):
                conservar_fin.append(hijo)
            elif hijo.tag == qn('w:r') and hijo.find(qn('w:commentReference')) is not None:
                conservar_fin.append(hijo)
            p.remove(hijo)
        for e in conservar_ini:
            p.append(e)
        for r in self._carreras(piezas):
            p.append(r)
        for e in conservar_fin:
            p.append(e)
        return p

    @staticmethod
    def _tam(elemento, defecto=18):
        for sz in elemento.iter(qn('w:sz')):
            return int(sz.get(qn('w:val')))
        return defecto

    def _llenar_celda(self, tc, piezas, tam):
        ps = tc.findall(qn('w:p'))
        ppr = ps[0].find(qn('w:pPr')) if ps else None
        ppr = copy.deepcopy(ppr) if ppr is not None else None
        for p in ps:
            tc.remove(p)
        partes = piezas if (isinstance(piezas, list) and piezas and isinstance(piezas[0], list)) else [piezas]
        for parte in partes:
            p = OxmlElement('w:p')
            if ppr is not None:
                p.append(copy.deepcopy(ppr))
            for r in self._carreras(parte, tam=tam):
                p.append(r)
            tc.append(p)

    def celda(self, tabla, fila, col, piezas):
        """Cambia el contenido de una celda (una lista de piezas, o una lista de párrafos), con su tamaño de letra."""
        tr = tabla.findall(qn('w:tr'))[fila]
        tc = tr.findall(qn('w:tc'))[col]
        self._llenar_celda(tc, piezas, self._tam(tc))

    def anadir_filas(self, tabla, filas, despues_de=None):
        """Añade filas copiando el formato de la última (o de la fila <despues_de>)."""
        trs = tabla.findall(qn('w:tr'))
        modelo = trs[-1] if despues_de is None else trs[despues_de]
        ancla = modelo
        nuevas = []
        for f in filas:
            tr = copy.deepcopy(modelo)
            ancla.addnext(tr)
            ancla = tr
            nuevas.append(tr)
            for tc, contenido in zip(tr.findall(qn('w:tc')), f):
                self._llenar_celda(tc, contenido, self._tam(tc))
        return nuevas

    def quitar_fila(self, tabla, indice):
        tr = tabla.findall(qn('w:tr'))[indice]
        tabla.remove(tr)

    @staticmethod
    def quitar(*elementos):
        for e in elementos:
            e.getparent().remove(e)

    def comentar(self, p, texto):
        """Un comentario sobre todo el párrafo <p>."""
        par = Paragraph(p, self.doc.part)
        runs = [r for r in par.runs if r._r.find(qn('w:commentReference')) is None and r.text]
        if not runs:
            runs = par.runs
        self.doc.add_comment(runs, text=texto, author=AUTOR, initials=INICIALES)

    # --- numeración ---

    def renumerar(self):
        """
        Recalcula el número que se ve en cada leyenda (campos SEQ Tabla y SEQ Figura) y el texto de
        cada referencia cruzada (campos REF _Ref_*), como haría Word al actualizar los campos.
        """
        cuenta = {'Tabla': 0, 'Figura': 0}
        destino = {}
        for p in self.body.iter(qn('w:p')):
            if estilo_de(p).startswith('TDC') or estilo_de(p) == 'Tabladeilustraciones':
                continue
            for instr, resultado in self._campos(p):
                m = re.match(r'\s*SEQ (Tabla|Figura) ', instr)
                if not m:
                    continue
                tipo = m.group(1)
                cuenta[tipo] += 1
                self._poner_resultado(resultado, str(cuenta[tipo]))
                for bs in p.iter(qn('w:bookmarkStart')):
                    if bs.get(qn('w:name')).startswith('_Ref_'):
                        destino[bs.get(qn('w:name'))] = f'{tipo} {cuenta[tipo]}'
        faltan = set()
        for p in self.body.iter(qn('w:p')):
            for instr, resultado in self._campos(p):
                m = re.match(r'\s*REF (_Ref_\S+)', instr)
                if not m:
                    continue
                nombre = m.group(1)
                if nombre in destino:
                    self._poner_resultado(resultado, destino[nombre])
                else:
                    faltan.add(nombre)
        return cuenta, faltan

    @staticmethod
    def _campos(p):
        """Pares (instrucción, carreras del resultado) de los campos simples del párrafo."""
        campos = []
        estado, instr, resultado = None, '', []
        for r in p.iter(qn('w:r')):
            fc = r.find(qn('w:fldChar'))
            if fc is not None:
                tipo = fc.get(qn('w:fldCharType'))
                if tipo == 'begin':
                    estado, instr, resultado = 'instr', '', []
                elif tipo == 'separate':
                    estado = 'res'
                elif tipo == 'end':
                    if estado is not None:
                        campos.append((instr, resultado))
                    estado = None
                continue
            if estado == 'instr':
                for it in r.iter(qn('w:instrText')):
                    instr += it.text or ''
            elif estado == 'res':
                resultado.append(r)
        return campos

    @staticmethod
    def _poner_resultado(carreras, texto):
        ts = [t for r in carreras for t in r.iter(qn('w:t'))]
        if not ts:
            return
        ts[0].text = texto
        for t in ts[1:]:
            t.text = ''

    def guardar(self, ruta):
        # Cada dibujo con su id propio (al quitar y añadir imágenes, python-docx puede repetirlos).
        wp = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
        for i, e in enumerate(self.doc.element.body.iter(wp + 'docPr'), start=1):
            e.set('id', str(i))
        # Que Word actualice al abrir los índices, los números de página y las referencias.
        ajustes = self.doc.settings.element
        if ajustes.find(qn('w:updateFields')) is None:
            u = OxmlElement('w:updateFields')
            u.set(qn('w:val'), 'true')
            # En su sitio del esquema: antes de hdrShapeDefaults, footnotePr, compat, rsids…
            despues = ('hdrShapeDefaults', 'footnotePr', 'endnotePr', 'compat', 'docVars', 'rsids', 'mathPr',
                       'attachedSchema', 'themeFontLang', 'clrSchemeMapping', 'doNotIncludeSubdocsInStats',
                       'doNotAutoCompressPictures', 'forceUpgrade', 'captions', 'readModeInkLockDown', 'smartTagType',
                       'schemaLibrary', 'shapeDefaults', 'doNotEmbedSmartTags', 'decimalSymbol', 'listSeparator')
            siguiente = next((h for h in ajustes if etree.QName(h).localname in despues), None)
            if siguiente is not None:
                siguiente.addprevious(u)
            else:
                ajustes.append(u)
        self.doc.save(ruta)
